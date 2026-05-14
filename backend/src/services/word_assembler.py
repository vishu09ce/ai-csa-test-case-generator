from datetime import datetime, timezone
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def _add_field_table(doc: Document, fields: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=len(fields), cols=2)
    table.style = "Table Grid"
    for i, (label, value) in enumerate(fields):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value or ""


def _add_data_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            table.rows[r_idx + 1].cells[c_idx].text = str(value) if value else ""


def _add_stamp(doc: Document, doc_code: str, project_id: str) -> None:
    timestamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    para = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    para.clear()
    run = para.add_run(
        f"AI GENERATED — UNALTERED  |  {doc_code}  |  Project ID: {project_id}  |  Generated: {timestamp}"
    )
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_signature_block(doc: Document, signatories: list[tuple[str, str]]) -> None:
    headers = ["Role", "Name", "Initials", "Date", "Action"]
    rows = [[role, "To be completed", "", "", action] for role, action in signatories]
    _add_data_table(doc, headers, rows)


def assemble_sap(data: dict, project_id: str, system_type: str) -> Document:
    doc = Document()
    _add_stamp(doc, "SAP", project_id)
    doc.add_heading("Software Assurance Plan", 0)
    doc.add_paragraph(f"Document 1 of 6  |  SAP  |  Template v1.1  |  FDA CSA Final Guidance (September 2025)")

    _add_heading(doc, "Section 1 — Document Control")
    _add_field_table(doc, [
        ("Document ID", f"SAP-{project_id}-001"),
        ("System Name", data.get("system_description", "")[:80]),
        ("Project ID", project_id),
        ("Document Number", "Document 1 of 6"),
        ("Template Version", "Template v1.1 (Auto-populated)"),
        ("Regulatory Framework", "FDA CSA Final Guidance, September 2025"),
        ("Document Status", "Draft"),
    ])

    _add_heading(doc, "Section 2 — System Identification")
    _add_field_table(doc, [
        ("System Description", data.get("system_description", "")),
        ("System Type", system_type),
        ("GAMP 5 Category", data.get("gamp5_category", "")),
        ("GAMP 5 Rationale", data.get("gamp5_rationale", "")),
        ("21 CFR Part 11 Applicability", str(data.get("part11_applicable", ""))),
        ("Part 11 Rationale", data.get("part11_rationale", "")),
    ])

    _add_heading(doc, "Section 3 — Intended Use Statements")
    ius = data.get("intended_uses", [])
    headers = ["ID", "Feature", "Intended Use Statement", "URS Ref", "FRS Ref", "BRD Ref", "Prod/QS", "Risk Tier"]
    rows = [[
        iu.get("id", ""), iu.get("feature", ""), iu.get("statement", ""),
        iu.get("urs_ref", ""), iu.get("frs_ref", ""), iu.get("brd_ref", ""),
        iu.get("prod_or_qs", ""), iu.get("risk_tier", "")
    ] for iu in ius]
    _add_data_table(doc, headers, rows)
    doc.add_paragraph("▶ HARD STOP — HITL GATE: PRA will not be generated until all signature blocks are completed.")

    _add_heading(doc, "Section 5 — Assurance Approach")
    doc.add_paragraph(data.get("assurance_strategy", ""))
    doc.add_paragraph(data.get("scripted_approach", ""))
    doc.add_paragraph(data.get("unscripted_approach", ""))

    _add_heading(doc, "Section 8 — Document Approval")
    _add_signature_block(doc, [
        ("Validation Lead", "Prepared and reviewed"),
        ("Reviewer", "Reviewed for accuracy and completeness"),
        ("Approver / System Owner", "Approved — authorizes PRA generation"),
    ])
    return doc


def assemble_pra(data: dict, project_id: str, system_type: str) -> Document:
    doc = Document()
    _add_stamp(doc, "PRA", project_id)
    doc.add_heading("Process Risk Assessment", 0)
    doc.add_paragraph(f"Document 2 of 6  |  PRA  |  Template v1.1  |  FDA CSA Final Guidance (September 2025)")

    _add_heading(doc, "Section 1 — Document Control")
    _add_field_table(doc, [
        ("Document ID", f"PRA-{project_id}-001"),
        ("Project ID", project_id),
        ("Document Number", "Document 2 of 6"),
        ("Template Version", "Template v1.1 (Auto-populated)"),
    ])

    _add_heading(doc, "Section 4 — Feature Risk Classification Matrix")
    classifications = data.get("risk_classifications", [])
    headers = ["ID", "Feature", "Failure Scenario", "Patient Safety?", "Product Quality?", "Risk Classification"]
    rows = [[
        r.get("id", ""), r.get("feature", ""), r.get("failure_scenario", ""),
        "Yes" if r.get("patient_safety_impact") else "No",
        "Yes" if r.get("product_quality_impact") else "No",
        r.get("risk_classification", "")
    ] for r in classifications]
    _add_data_table(doc, headers, rows)

    _add_heading(doc, "Section 5 — Risk Classification Rationale")
    headers = ["ID", "Feature", "Classification", "Rationale"]
    rows = [[
        r.get("id", ""), r.get("feature", ""),
        r.get("risk_classification", ""), r.get("rationale", "")
    ] for r in classifications]
    _add_data_table(doc, headers, rows)

    _add_heading(doc, "Section 6 — Risk Classification Summary")
    _add_field_table(doc, [
        ("Total features assessed", str(data.get("total_features", len(classifications)))),
        ("High Process Risk", str(data.get("high_risk_count", ""))),
        ("Not High Process Risk", str(data.get("not_high_risk_count", ""))),
    ])
    doc.add_paragraph("▶ HARD STOP — HITL GATE: RTM, Scripted Protocol, and Unscripted Record will not be generated until all signature blocks are completed.")

    _add_heading(doc, "Section 7 — Document Approval")
    _add_signature_block(doc, [
        ("Reviewer", "Reviewed — confirms all risk classifications"),
        ("Approver", "Approved — final risk determination; authorizes test generation"),
        ("System Owner", "Acknowledged — confirms organizational context"),
    ])
    return doc


def assemble_rtm(data: dict, project_id: str) -> Document:
    doc = Document()
    _add_stamp(doc, "RTM", project_id)
    doc.add_heading("Requirements Traceability Matrix", 0)
    doc.add_paragraph(f"Document 3 of 6  |  RTM  |  Template v1.1  |  FDA CSA Final Guidance (September 2025)")

    _add_heading(doc, "Section 1 — Document Control")
    _add_field_table(doc, [
        ("Document ID", f"RTM-{project_id}-001"),
        ("Project ID", project_id),
        ("Document Number", "Document 3 of 6"),
        ("Template Version", "Template v1.1 (Auto-populated)"),
    ])

    _add_heading(doc, "Section 3 — Traceability Matrix")
    requirements = data.get("requirements", [])
    headers = ["Req ID", "Requirement Summary", "URS Ref", "FRS Ref", "BRD Ref", "IU Ref", "Risk Class", "Test Case ID", "Result"]
    rows = [[
        r.get("req_id", ""), r.get("summary", ""), r.get("urs_ref", ""),
        r.get("frs_ref", ""), r.get("brd_ref", ""), r.get("iu_ref", ""),
        r.get("risk_classification", ""), r.get("test_case_id", ""), r.get("result", "")
    ] for r in requirements]
    _add_data_table(doc, headers, rows)

    _add_heading(doc, "Section 5 — Reviewer Sign-Off")
    _add_signature_block(doc, [
        ("Reviewer", "Reviewed — confirms RTM completeness and traceability accuracy"),
        ("Validation Lead", "Acknowledged — confirms alignment with Documents 1 and 2"),
    ])
    return doc


def assemble_stp(data: dict, project_id: str) -> Document:
    doc = Document()
    _add_stamp(doc, "STP", project_id)
    doc.add_heading("Scripted Test Protocol and Execution Record", 0)
    doc.add_paragraph(f"Document 4 of 6  |  STP  |  Template v1.1  |  FDA CSA Final Guidance (September 2025)")

    _add_heading(doc, "Section 1 — Document Control")
    _add_field_table(doc, [
        ("Document ID", f"STP-{project_id}-001"),
        ("Project ID", project_id),
        ("Document Number", "Document 4 of 6"),
        ("Template Version", "Template v1.1 (Auto-populated)"),
    ])

    _add_heading(doc, "Section 4 — Test Cases")
    test_cases = data.get("test_cases", [])
    headers = ["TC ID", "Req ID", "Preconditions", "Test Steps", "Expected Result", "Actual Result", "Pass/Fail"]
    rows = [[
        tc.get("tc_id", ""), tc.get("req_id", ""), tc.get("preconditions", ""),
        tc.get("test_steps", ""), tc.get("expected_result", ""),
        tc.get("actual_result", ""), tc.get("pass_fail", "")
    ] for tc in test_cases]
    _add_data_table(doc, headers, rows)

    doc.add_paragraph("▶ HARD STOP — HITL GATE: No test execution begins without a completed Approver signature.")

    _add_heading(doc, "Section 6 — Deviation Log")
    _add_data_table(doc, ["Dev ID", "TC ID", "Description", "Resolution", "Root Cause", "Status", "Resolved By"], [])

    _add_heading(doc, "Section 7 — Protocol Approval (Pre-Execution)")
    _add_signature_block(doc, [
        ("Reviewer", "Reviewed — confirms test cases are accurate and complete"),
        ("Approver", "APPROVED FOR EXECUTION — authorizes execution to begin"),
    ])

    _add_heading(doc, "Section 8 — Execution Summary and Record Closure")
    _add_signature_block(doc, [
        ("Executor", "Confirms testing was performed as documented"),
        ("Approver", "RECORD CLOSED — confirms high-risk features are fit for intended use"),
    ])
    return doc


def assemble_utr(data: dict, project_id: str) -> Document:
    doc = Document()
    _add_stamp(doc, "UTR", project_id)
    doc.add_heading("Unscripted Test Record and Execution Record", 0)
    doc.add_paragraph(f"Document 5 of 6  |  UTR  |  Template v1.1  |  FDA CSA Final Guidance (September 2025)")

    _add_heading(doc, "Section 1 — Document Control")
    _add_field_table(doc, [
        ("Document ID", f"UTR-{project_id}-001"),
        ("Project ID", project_id),
        ("Document Number", "Document 5 of 6"),
        ("Template Version", "Template v1.1 (Auto-populated)"),
    ])

    _add_heading(doc, "Section 4 — Test Records")
    test_records = data.get("test_records", [])
    headers = ["UTR ID", "Req ID", "Feature", "AI-Suggested Exploratory Scenarios", "Tester Observations", "Conclusion"]
    rows = [[
        tr.get("utr_id", ""), tr.get("req_id", ""), tr.get("feature", ""),
        tr.get("exploratory_scenarios", ""), tr.get("tester_observations", ""),
        tr.get("conclusion", "")
    ] for tr in test_records]
    _add_data_table(doc, headers, rows)

    _add_heading(doc, "Section 6 — Record Closure")
    _add_signature_block(doc, [
        ("Executor", "Confirms exploratory testing was conducted and observations are accurately recorded"),
        ("Approver", "RECORD CLOSED — confirms not-high-risk features are fit for intended use"),
    ])
    return doc


def assemble_asr(data: dict, project_id: str) -> Document:
    doc = Document()
    _add_stamp(doc, "ASR", project_id)
    doc.add_heading("Assurance Summary Report", 0)
    doc.add_paragraph(f"Document 6 of 6  |  ASR  |  Template v1.1  |  FDA CSA Final Guidance (September 2025)")

    _add_heading(doc, "Section 1 — Document Control")
    _add_field_table(doc, [
        ("Document ID", f"ASR-{project_id}-001"),
        ("Project ID", project_id),
        ("Document Number", "Document 6 of 6"),
        ("Template Version", "Template v1.1 (Auto-populated)"),
    ])

    _add_heading(doc, "Section 3 — Intended Use Coverage Confirmation")
    coverage = data.get("coverage_summary", [])
    headers = ["IU ID", "Feature", "Risk Class", "Test Doc", "Result", "Conclusion"]
    rows = [[
        c.get("iu_id", ""), c.get("feature", ""), c.get("risk_class", ""),
        c.get("test_doc", ""), c.get("result", ""), c.get("conclusion", "")
    ] for c in coverage]
    _add_data_table(doc, headers, rows)

    _add_heading(doc, "Section 6 — Overall Assurance Conclusion")
    doc.add_paragraph(data.get("overall_conclusion", ""))
    _add_field_table(doc, [
        ("System fitness for intended use", data.get("system_fitness", "")),
        ("All intended uses addressed", str(data.get("all_intended_uses_addressed", ""))),
        ("All risk classifications confirmed", str(data.get("all_risk_classifications_confirmed", ""))),
    ])
    doc.add_paragraph("▶ HARD STOP — HITL GATE: Terminal gate. Cannot be signed until all upstream documents are Approved or Closed.")

    _add_heading(doc, "Section 8 — Final Approval")
    _add_signature_block(doc, [
        ("Validation Lead", "Confirms all assurance activities were conducted as documented"),
        ("Reviewer", "Confirms the evidence package is complete and accurate"),
        ("Approver / System Owner", "APPROVED — declares system fit for intended use"),
    ])
    return doc
