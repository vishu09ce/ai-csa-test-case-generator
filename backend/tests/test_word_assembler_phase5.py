import pytest
from docx import Document
from docx.document import Document as DocxDocument
from backend.src.services.word_assembler import assemble_stp, assemble_utr


def _find_table_with_header(doc: Document, header_text: str):
    for table in doc.tables:
        if any(header_text in cell.text for cell in table.rows[0].cells):
            return table
    return None


SAMPLE_TC_HIGH = {
    "tc_id": "TC-001",
    "req_id": "REQ-001",
    "preconditions": ["User is logged in with Admin role"],
    "test_steps": ["Navigate to page", "Enter value", "Click Submit"],
    "expected_result": "System displays confirmation.",
    "actual_result": "",
    "pass_fail": "",
    "confidence": "High",
    "flags": [],
}

SAMPLE_TC_LOW = {
    **SAMPLE_TC_HIGH,
    "tc_id": "TC-002",
    "confidence": "Low",
    "flags": ["VALIDATION_FAILED"],
}

SAMPLE_TC_MEDIUM = {
    **SAMPLE_TC_HIGH,
    "tc_id": "TC-003",
    "confidence": "Medium",
    "flags": ["INFERRED_FRS_LINK"],
}

STP_DATA = {"test_cases": [SAMPLE_TC_HIGH, SAMPLE_TC_LOW, SAMPLE_TC_MEDIUM]}

SAMPLE_UTR = {
    "utr_id": "UTR-001",
    "req_id": "REQ-001",
    "feature_description": "Dashboard summary view.",
    "exploratory_scenarios": [
        "Normal use: verify dashboard loads.",
        "Boundary: verify with zero records.",
        "Error state: verify with network disconnected.",
    ],
    "tester_observations": "",
    "conclusion": "",
    "confidence": "Medium",
    "flags": ["INFERRED_FRS_LINK"],
}

UTR_DATA = {"test_records": [SAMPLE_UTR]}


class TestAssembleStp:
    def test_table_has_confidence_column(self):
        doc = assemble_stp(STP_DATA, "PROJ-001")
        assert _find_table_with_header(doc, "AI Confidence") is not None

    def test_table_has_flags_column(self):
        doc = assemble_stp(STP_DATA, "PROJ-001")
        assert _find_table_with_header(doc, "Flags") is not None

    def test_confidence_values_written_in_rows(self):
        doc = assemble_stp(STP_DATA, "PROJ-001")
        table = _find_table_with_header(doc, "AI Confidence")
        cell_texts = [cell.text for row in table.rows[1:] for cell in row.cells]
        assert "High" in cell_texts
        assert "Medium" in cell_texts
        assert "Low" in cell_texts

    def test_validation_failed_flag_written(self):
        doc = assemble_stp(STP_DATA, "PROJ-001")
        table = _find_table_with_header(doc, "Flags")
        cell_texts = [cell.text for row in table.rows[1:] for cell in row.cells]
        assert any("VALIDATION_FAILED" in t for t in cell_texts)

    def test_preconditions_list_joined(self):
        doc = assemble_stp(STP_DATA, "PROJ-001")
        table = _find_table_with_header(doc, "Preconditions")
        cell_texts = [cell.text for row in table.rows[1:] for cell in row.cells]
        assert any("User is logged in with Admin role" in t for t in cell_texts)

    def test_confidence_distribution_table_present(self):
        doc = assemble_stp(STP_DATA, "PROJ-001")
        assert _find_table_with_header(doc, "Confidence Level") is not None

    def test_confidence_distribution_counts_correct(self):
        doc = assemble_stp(STP_DATA, "PROJ-001")
        table = _find_table_with_header(doc, "Confidence Level")
        rows = {row.cells[0].text: row.cells[1].text for row in table.rows[1:]}
        assert rows["High"] == "1"
        assert rows["Medium"] == "1"
        assert rows["Low"] == "1"
        assert rows["VALIDATION_FAILED"] == "1"

    def test_template_version_is_v12(self):
        doc = assemble_stp(STP_DATA, "PROJ-001")
        texts = [p.text for p in doc.paragraphs]
        assert any("v1.2" in t for t in texts)

    def test_empty_test_cases_does_not_raise(self):
        doc = assemble_stp({"test_cases": []}, "PROJ-001")
        assert isinstance(doc, DocxDocument)

    def test_empty_distribution_counts_are_zero(self):
        doc = assemble_stp({"test_cases": []}, "PROJ-001")
        table = _find_table_with_header(doc, "Confidence Level")
        rows = {row.cells[0].text: row.cells[1].text for row in table.rows[1:]}
        assert rows["High"] == "0"
        assert rows["VALIDATION_FAILED"] == "0"


class TestAssembleUtr:
    def test_table_has_confidence_column(self):
        doc = assemble_utr(UTR_DATA, "PROJ-001")
        assert _find_table_with_header(doc, "AI Confidence") is not None

    def test_table_has_flags_column(self):
        doc = assemble_utr(UTR_DATA, "PROJ-001")
        assert _find_table_with_header(doc, "Flags") is not None

    def test_feature_description_field_used(self):
        doc = assemble_utr(UTR_DATA, "PROJ-001")
        table = _find_table_with_header(doc, "Feature Description")
        cell_texts = [cell.text for row in table.rows[1:] for cell in row.cells]
        assert any("Dashboard summary view." in t for t in cell_texts)

    def test_exploratory_scenarios_list_joined(self):
        doc = assemble_utr(UTR_DATA, "PROJ-001")
        table = _find_table_with_header(doc, "AI-Suggested Exploratory Scenarios")
        cell_texts = [cell.text for row in table.rows[1:] for cell in row.cells]
        assert any("Normal use" in t for t in cell_texts)

    def test_confidence_distribution_table_present(self):
        doc = assemble_utr(UTR_DATA, "PROJ-001")
        assert _find_table_with_header(doc, "Confidence Level") is not None

    def test_confidence_distribution_counts_correct(self):
        doc = assemble_utr(UTR_DATA, "PROJ-001")
        table = _find_table_with_header(doc, "Confidence Level")
        rows = {row.cells[0].text: row.cells[1].text for row in table.rows[1:]}
        assert rows["Medium"] == "1"
        assert rows["High"] == "0"
        assert rows["VALIDATION_FAILED"] == "0"

    def test_template_version_is_v12(self):
        doc = assemble_utr(UTR_DATA, "PROJ-001")
        texts = [p.text for p in doc.paragraphs]
        assert any("v1.2" in t for t in texts)

    def test_empty_test_records_does_not_raise(self):
        doc = assemble_utr({"test_records": []}, "PROJ-001")
        assert isinstance(doc, DocxDocument)
